# Builds the next Merchant Center feed workbook from the previous one plus the current outlet data.
# 0) node --env-file=.env.local scripts/outlet/sync-stock.mjs   (refresh the live stock snapshot first)
# 1) npx tsx scripts/outlet/dump-feed-data.ts feed-data.json
# 2) python scripts/outlet/build-merchant-feed.py <previous.xlsx> feed-data.json <new.xlsx>
# Existing rows keep their text (titles are Stefan's); only price, link, image, availability and the
# reference quantity are refreshed. Rows for newly eligible outlet SKUs are appended. Quantity is the
# live stock (quantity_remaining) and a row at zero becomes "out of stock". Every row gets
# google_product_category 503729 (Wire Terminals & Connectors, Stefan's Sheet value); the column is added
# after mpn when the previous workbook lacks it. Requires openpyxl.
import copy, csv, json, re, sys
from openpyxl import load_workbook

prev_path, data_path, out_path = sys.argv[1:4]
TAIL = "Surplus stock from Adcontact's own warehouse, sold at outlet pricing while quantities last."
CATEGORY_NOUN = {"Accessories": "Accessory", "Contacts": "Contact", "Tools": "Tool", "Connectors": "Connector"}
CONNECTOR_PREFIXES = ("IMC", "WT", "DT16", "HDP24", "8N1534")
GOOGLE_PRODUCT_CATEGORY = 503729

data = json.load(open(data_path, encoding="utf8"))
wb = load_workbook(prev_path)
ws = wb.worksheets[0]
header = [c.value for c in ws[1]]
if "google_product_category" not in header:
    at = header.index("mpn") + 2  # 1-based column right after mpn
    ws.insert_cols(at)
    for r in range(1, ws.max_row + 1):
        src = ws.cell(r, at - 1)
        if src.has_style:
            ws.cell(r, at)._style = copy.copy(src._style)
    ws.cell(1, at, "google_product_category")
    header = [c.value for c in ws[1]]
col = {h: i for i, h in enumerate(header)}
norm = lambda s: re.sub(r"[^a-z0-9]", "", s.lower())

def split_note(description, mpn):
    d = re.sub(r"\s+hand tool$", " Hand tool", description.strip(), flags=re.I)
    for i in range(1, len(d) + 1):
        if norm(d[:i]) == norm(mpn):
            return d[i:].strip()
    return ""

def new_row(r):
    if r["kind"] == "deutsch":
        title = re.sub(r",\s*connector,\s*for wire processing$", "", r["seoTitle"] or f"Deutsch {r['partNumber']}")
        mpn = r["partNumber"]
        parts = title.split(", ")[1:]
        m = re.match(r"(\d+)-Way (\w+)$", parts[-1]) if len(parts) > 1 else None
        if m:
            desc = f"{mpn}, {parts[0]}, {m.group(1)}-way, {m.group(2).lower()} connector from Deutsch. {TAIL}"
        elif parts:
            desc = f"{mpn}, {parts[0]} connector from Deutsch. {TAIL}"
        else:
            desc = f"{mpn} connector from Deutsch. {TAIL}"
    elif r.get("seo"):
        mpn = r["partNumber"]
        title = r["seo"]["shortTitle"]
        desc = r["seo"]["feedDescription"]
    elif r["kind"] == "catalogue":
        mpn = r["partNumber"]
        noun = CATEGORY_NOUN.get(r["categoryName"] or "", "")
        note = split_note(r["description"], mpn)
        if not note:
            note = noun
        elif len(note) <= 3 and noun:
            note = f"{noun}, {note}"
        title = f"Deutsch {mpn}" + (f", {note}" if note else "")
        if not note:
            phrase = mpn
        elif re.search(r"\d", note):
            phrase = f"{mpn}, {note}"
        else:
            phrase = f"{mpn} {note[0].lower() + note[1:]}"
        desc = f"{phrase} from Deutsch. {TAIL}"
    else:
        mpn = r["partNumber"]
        noun = "connector" if mpn.upper().startswith(CONNECTOR_PREFIXES) else ""
        title = f"Deutsch {mpn}" + (f", {noun}" if noun else "")
        desc = f"{mpn}{' ' + noun if noun else ''} from Deutsch. {TAIL}"
    if r["reference"]:
        desc += " Image shows a similar part."
    return {"id": r["sku"], "title": title, "description": desc, "link": r["link"], "image_link": r["imageLink"],
            "availability": "in stock" if r["quantity"] > 0 else "out of stock", "price": f"{r['priceEur']:.2f} EUR", "condition": "new", "brand": "Deutsch",
            "mpn": mpn, "google_product_category": GOOGLE_PRODUCT_CATEGORY, "qty_available_reference_only": r["quantity"]}

by_id = {r["sku"]: r for r in data}
seen, drift = set(), []
for row in ws.iter_rows(min_row=2):
    rid = str(row[col["id"]].value or "")
    if not rid: continue
    seen.add(rid)
    cat_cell = row[col["google_product_category"]]
    if cat_cell.value in (None, ""):
        cat_cell.value = GOOGLE_PRODUCT_CATEGORY
        drift.append((rid, "google_product_category set to 503729"))
    r = by_id.get(rid)
    if r is None:
        drift.append((rid, "no longer eligible")); continue
    want = {"price": f"{r['priceEur']:.2f} EUR", "link": r["link"], "image_link": r["imageLink"], "qty_available_reference_only": r["quantity"],
            "availability": "in stock" if r["quantity"] > 0 else "out of stock"}
    for k, v in want.items():
        cell = row[col[k]]
        if str(cell.value) != str(v):
            drift.append((rid, f"{k}: {cell.value} -> {v}")); cell.value = v

last = ws.max_row
while last > 1 and not ws.cell(last, 1).value: last -= 1
added = []
for r in data:
    if r["sku"] in seen: continue
    values = new_row(r)
    last += 1
    for h, i in col.items():
        cell = ws.cell(last, i + 1, values.get(h))
        src = ws.cell(last - 1, i + 1)
        if src.has_style:
            cell._style = copy.copy(src._style)
    added.append(values)

wb.save(out_path)
with open(re.sub(r"\.xlsx$", ".csv", out_path), "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f); w.writerow(header)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]: w.writerow(row)
print(f"existing rows kept: {len(seen)} | refreshed/removed (drift): {len(drift)} | new rows added: {len(added)} | total: {len(seen) + len(added)}")
for d in drift[:20]: print("  drift:", d)
json.dump(added, open(re.sub(r"\.xlsx$", ".new-rows.json", out_path), "w", encoding="utf8"), indent=1)
